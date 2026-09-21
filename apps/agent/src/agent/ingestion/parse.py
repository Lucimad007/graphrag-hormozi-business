from __future__ import annotations

import os
import re
import zipfile
from io import BytesIO
from pathlib import Path

from pypdf import PdfReader

from agent.ingestion.models import ParsedDocument, ParsedSection
from agent.ingestion.normalize import document_id_for, normalize_text

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_HTML_TAG = re.compile(r"<[^>]+>")

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx", ".epub"}


class UnsupportedDocumentError(ValueError):
    pass


def parse_path(path: Path, *, document_id: str | None = None) -> ParsedDocument:
    raw = path.read_bytes()
    return parse_bytes(
        raw,
        source=str(path),
        suffix=path.suffix,
        document_id=document_id,
    )


def parse_bytes(
    raw: bytes,
    *,
    source: str,
    suffix: str,
    document_id: str | None = None,
) -> ParsedDocument:
    ext = suffix.lower()
    doc_id = document_id or document_id_for(source, raw)
    if ext in {".txt", ".md", ".markdown"}:
        text = raw.decode("utf-8")
        media = "text/markdown" if ext in {".md", ".markdown"} else "text/plain"
        sections = (
            _sections_from_markdown(text) if ext in {".md", ".markdown"} else _single_section(text)
        )
        return ParsedDocument(
            document_id=doc_id,
            source=source,
            media_type=media,
            sections=sections,
            metadata={"filename": Path(source).name},
        )
    if ext == ".pdf":
        return _parse_pdf(raw, source=source, document_id=doc_id)
    if ext == ".docx":
        return _parse_docx(raw, source=source, document_id=doc_id)
    if ext == ".epub":
        return _parse_epub(raw, source=source, document_id=doc_id)
    raise UnsupportedDocumentError(f"Unsupported document type: {suffix}")


def _single_section(text: str) -> list[ParsedSection]:
    normalized = normalize_text(text)
    if not normalized:
        return []
    return [ParsedSection(title=None, text=normalized)]


def _sections_from_markdown(text: str) -> list[ParsedSection]:
    sections: list[ParsedSection] = []
    current_title: str | None = None
    buf: list[str] = []

    def flush() -> None:
        body = normalize_text("\n".join(buf))
        if body:
            sections.append(ParsedSection(title=current_title, text=body))
        buf.clear()

    for line in text.replace("\r\n", "\n").split("\n"):
        match = _HEADING.match(line.strip())
        if match:
            flush()
            current_title = normalize_text(match.group(2))
            continue
        buf.append(line)
    flush()
    return sections


def _ocr_enabled() -> bool:
    flag = os.environ.get("INGEST_PDF_OCR", "true").strip().lower()
    return flag not in {"0", "false", "no"}


def _pdf_pypdf_sections(raw: bytes) -> tuple[list[ParsedSection], int]:
    reader = PdfReader(BytesIO(raw))
    sections: list[ParsedSection] = []
    for index, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        normalized = normalize_text(extracted)
        if not normalized:
            continue
        sections.append(ParsedSection(title=f"page {index}", text=normalized))
    return sections, len(reader.pages)


def _pdf_pymupdf_sections(raw: bytes) -> list[ParsedSection]:
    try:
        import pymupdf
    except ImportError:
        return []
    try:
        document = pymupdf.open(stream=raw, filetype="pdf")
    except Exception:
        return []
    try:
        sections: list[ParsedSection] = []
        for index, page in enumerate(document, start=1):
            normalized = normalize_text(page.get_text("text") or "")
            if normalized:
                sections.append(ParsedSection(title=f"page {index}", text=normalized))
        return sections
    finally:
        document.close()


_OCR_ENGINE = None


def _ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR

        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _pdf_ocr_sections(raw: bytes) -> list[ParsedSection]:
    try:
        import numpy as np
        import pymupdf
    except ImportError:
        return []
    try:
        document = pymupdf.open(stream=raw, filetype="pdf")
    except Exception:
        return []
    engine = _ocr_engine()
    try:
        sections: list[ParsedSection] = []
        for index, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
                pixmap.height, pixmap.width, pixmap.n
            )
            result, _elapsed = engine(image)
            lines: list[str] = []
            for item in result or []:
                if len(item) >= 2 and item[1]:
                    lines.append(str(item[1]))
            normalized = normalize_text("\n".join(lines))
            if normalized:
                sections.append(ParsedSection(title=f"page {index}", text=normalized))
        return sections
    finally:
        document.close()


def _parse_pdf(raw: bytes, *, source: str, document_id: str) -> ParsedDocument:
    sections, page_count = _pdf_pypdf_sections(raw)
    used_ocr = False
    if not sections:
        sections = _pdf_pymupdf_sections(raw)
    if not sections and _ocr_enabled():
        sections = _pdf_ocr_sections(raw)
        used_ocr = bool(sections)
    return ParsedDocument(
        document_id=document_id,
        source=source,
        media_type="application/pdf",
        sections=sections,
        metadata={
            "filename": Path(source).name,
            "page_count": page_count,
            "ocr": used_ocr,
        },
    )


def _parse_docx(raw: bytes, *, source: str, document_id: str) -> ParsedDocument:
    from docx import Document

    document = Document(BytesIO(raw))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return ParsedDocument(
        document_id=document_id,
        source=source,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sections=_single_section(text),
        metadata={"filename": Path(source).name},
    )


def _parse_epub(raw: bytes, *, source: str, document_id: str) -> ParsedDocument:
    parts: list[str] = []
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        names = [
            name
            for name in archive.namelist()
            if name.lower().endswith((".xhtml", ".html", ".htm"))
        ]
        for name in sorted(names):
            html = archive.read(name).decode("utf-8", errors="ignore")
            parts.append(_HTML_TAG.sub(" ", html))
    return ParsedDocument(
        document_id=document_id,
        source=source,
        media_type="application/epub+zip",
        sections=_single_section("\n".join(parts)),
        metadata={"filename": Path(source).name},
    )


def iter_ingestible_files(
    root: Path, suffixes: set[str] | None = None
) -> list[Path]:
    allowed = {item.lower() for item in (suffixes or TEXT_SUFFIXES)}
    files = [
        path
        for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in allowed
    ]
    return sorted(files)
